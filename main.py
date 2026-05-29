import ctypes
import json
import logging
import os
import sqlite3
import subprocess
import sys
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from PyQt6.QtCore import QCalendar, QDateTime, QTimer, Qt
from PyQt6.QtGui import QAction, QCloseEvent, QCursor, QFont, QFontDatabase, QIcon
from PyQt6.QtNetwork import QLocalServer, QLocalSocket
from PyQt6.QtWidgets import (
    QApplication,
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDateEdit,
    QDateTimeEdit,
    QDialog,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMenu,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QStyle,
    QSystemTrayIcon,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

APP_NAME = "TimeIncomeTracker"
SINGLE_INSTANCE_KEY = f"{APP_NAME}.single-instance"
PREFERRED_FONT_FAMILIES = ["IRANSans", "IRANSansX", "Iran Sans", "Yekan", "B Yekan", "Vazirmatn", "Tahoma", "Segoe UI"]
APP_FONT_SIZE = 10
AUTO_RESUME_ACTIVITY_SECONDS = 5


PERSIAN_CALENDAR = QCalendar(QCalendar.System.Jalali)


def configure_persian_date_edit(widget):
    widget.setCalendar(PERSIAN_CALENDAR)
    widget.setDisplayFormat("yyyy/MM/dd")


def configure_persian_datetime_edit(widget):
    widget.setCalendar(PERSIAN_CALENDAR)
    widget.setDisplayFormat("yyyy/MM/dd HH:mm:ss")


def gregorian_to_jalali(year: int, month: int, day: int):
    g_days_in_month = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    j_days_in_month = [31, 31, 31, 31, 31, 31, 30, 30, 30, 30, 30, 29]
    gy = year - 1600
    gm = month - 1
    gd = day - 1
    g_day_no = 365 * gy + (gy + 3) // 4 - (gy + 99) // 100 + (gy + 399) // 400
    for i in range(gm):
        g_day_no += g_days_in_month[i]
    if gm > 1 and ((gy + 1600) % 4 == 0 and ((gy + 1600) % 100 != 0 or (gy + 1600) % 400 == 0)):
        g_day_no += 1
    g_day_no += gd
    j_day_no = g_day_no - 79
    j_np = j_day_no // 12053
    j_day_no %= 12053
    jy = 979 + 33 * j_np + 4 * (j_day_no // 1461)
    j_day_no %= 1461
    if j_day_no >= 366:
        jy += (j_day_no - 1) // 365
        j_day_no = (j_day_no - 1) % 365
    for i in range(11):
        if j_day_no < j_days_in_month[i]:
            break
        j_day_no -= j_days_in_month[i]
    return jy, i + 1, j_day_no + 1



def jalali_to_gregorian(year: int, month: int, day: int):
    g_days_in_month = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    j_days_in_month = [31, 31, 31, 31, 31, 31, 30, 30, 30, 30, 30, 29]
    jy = year - 979
    jm = month - 1
    jd = day - 1
    j_day_no = 365 * jy + (jy // 33) * 8 + ((jy % 33) + 3) // 4
    for i in range(jm):
        j_day_no += j_days_in_month[i]
    j_day_no += jd
    g_day_no = j_day_no + 79
    gy = 1600 + 400 * (g_day_no // 146097)
    g_day_no %= 146097
    leap = True
    if g_day_no >= 36525:
        g_day_no -= 1
        gy += 100 * (g_day_no // 36524)
        g_day_no %= 36524
        if g_day_no >= 365:
            g_day_no += 1
        else:
            leap = False
    gy += 4 * (g_day_no // 1461)
    g_day_no %= 1461
    if g_day_no >= 366:
        leap = False
        g_day_no -= 1
        gy += g_day_no // 365
        g_day_no %= 365
    for i in range(12):
        days = g_days_in_month[i] + (1 if i == 1 and leap else 0)
        if g_day_no < days:
            break
        g_day_no -= days
    return gy, i + 1, g_day_no + 1


def jalali_month_range(value):
    jy, jm, _ = gregorian_to_jalali(value.year, value.month, value.day)
    start = datetime(*jalali_to_gregorian(jy, jm, 1)).date()
    if jm == 12:
        next_jy, next_jm = jy + 1, 1
    else:
        next_jy, next_jm = jy, jm + 1
    next_month = datetime(*jalali_to_gregorian(next_jy, next_jm, 1)).date()
    return start, next_month - timedelta(days=1)

def format_jalali_date(value) -> str:
    if isinstance(value, datetime):
        value = value.date()
    jy, jm, jd = gregorian_to_jalali(value.year, value.month, value.day)
    return f"{jy:04}/{jm:02}/{jd:02}"


def preferred_font_family() -> str:
    installed = set(QFontDatabase.families())
    for family in PREFERRED_FONT_FAMILIES:
        if family in installed:
            return family
    return PREFERRED_FONT_FAMILIES[-1]


def apply_app_font(app: QApplication) -> str:
    family = preferred_font_family()
    app.setFont(QFont(family, APP_FONT_SIZE))
    app.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
    return family


def stylesheet_font_family() -> str:
    return ", ".join(f'"{family}"' for family in PREFERRED_FONT_FAMILIES)


RTL_ALIGNMENT = Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter


def configure_rtl_widget(widget: QWidget):
    widget.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
    for child in widget.findChildren(QWidget):
        child.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        if isinstance(child, (QLineEdit, QSpinBox, QDateEdit, QDateTimeEdit)):
            child.setAlignment(RTL_ALIGNMENT)
        elif isinstance(child, QTextEdit):
            child.setAlignment(Qt.AlignmentFlag.AlignRight)
        elif isinstance(child, QLabel):
            child.setAlignment(RTL_ALIGNMENT)
        elif isinstance(child, QTableWidget):
            child.horizontalHeader().setDefaultAlignment(RTL_ALIGNMENT)
            child.verticalHeader().setDefaultAlignment(RTL_ALIGNMENT)


def rtl_item(value) -> QTableWidgetItem:
    item = QTableWidgetItem(str(value))
    item.setTextAlignment(RTL_ALIGNMENT)
    return item


def app_data_dir() -> Path:
    base = Path(os.environ.get("APPDATA", str(Path.home())))
    path = base / APP_NAME
    path.mkdir(parents=True, exist_ok=True)
    (path / "logs").mkdir(exist_ok=True)
    return path


def setup_logging(base_path: Path):
    log_file = base_path / "logs" / "app.log"
    logging.basicConfig(level=logging.INFO, filename=log_file, format="%(asctime)s %(levelname)s %(message)s", encoding="utf-8")


class DB:
    def __init__(self, db_path: Path):
        self.conn = sqlite3.connect(db_path)
        self.conn.row_factory = sqlite3.Row
        self.init()

    def init(self):
        cur = self.conn.cursor()
        cur.executescript(
            """
            CREATE TABLE IF NOT EXISTS projects (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                hourly_rate INTEGER NULL,
                has_hourly_rate INTEGER DEFAULT 0,
                is_active INTEGER DEFAULT 1,
                created_at TEXT,
                updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS time_entries (
                id INTEGER PRIMARY KEY,
                project_id INTEGER,
                project_name_snapshot TEXT,
                task_description TEXT,
                start_time TEXT,
                end_time TEXT,
                duration_seconds INTEGER,
                hourly_rate_snapshot INTEGER NULL,
                amount INTEGER NULL,
                status TEXT,
                created_at TEXT,
                updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT
            );
            CREATE TABLE IF NOT EXISTS active_timer_state (
                id INTEGER PRIMARY KEY,
                project_id INTEGER,
                project_name_snapshot TEXT,
                task_description TEXT,
                start_time TEXT,
                accumulated_seconds INTEGER,
                is_paused INTEGER,
                pause_started_at TEXT,
                auto_paused INTEGER,
                hourly_rate_snapshot INTEGER NULL,
                updated_at TEXT
            );
            """
        )
        self.conn.commit()

    def set_setting(self, key, value):
        self.conn.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, str(value)))
        self.conn.commit()

    def get_setting(self, key, default=None):
        row = self.conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
        return row[0] if row else default


class SingleInstanceServer:
    def __init__(self, show_callback):
        self.show_callback = show_callback
        self.server = QLocalServer()
        self.server.newConnection.connect(self.on_new_connection)

    @staticmethod
    def signal_existing() -> bool:
        socket = QLocalSocket()
        socket.connectToServer(SINGLE_INSTANCE_KEY)
        if socket.waitForConnected(250):
            socket.write(b"show")
            socket.flush()
            socket.waitForBytesWritten(500)
            socket.disconnectFromServer()
            return True
        return False

    def listen(self) -> bool:
        QLocalServer.removeServer(SINGLE_INSTANCE_KEY)
        return self.server.listen(SINGLE_INSTANCE_KEY)

    def on_new_connection(self):
        while self.server.hasPendingConnections():
            socket = self.server.nextPendingConnection()
            socket.readyRead.connect(socket.readAll)
            socket.disconnected.connect(socket.deleteLater)
            self.show_callback()


class IdleMonitor:
    class LASTINPUTINFO(ctypes.Structure):
        _fields_ = [("cbSize", ctypes.c_uint), ("dwTime", ctypes.c_uint)]

    @staticmethod
    def idle_seconds() -> int:
        if os.name != "nt":
            return 0
        user32 = ctypes.windll.user32
        kernel32 = ctypes.windll.kernel32
        lii = IdleMonitor.LASTINPUTINFO()
        lii.cbSize = ctypes.sizeof(lii)
        if not user32.GetLastInputInfo(ctypes.byref(lii)):
            return 0
        tick = kernel32.GetTickCount()
        return int((tick - lii.dwTime) / 1000)


@dataclass
class TimerState:
    project_id: int
    project_name: str
    task_description: str
    start_time: datetime
    accumulated_seconds: int = 0
    is_paused: bool = False
    pause_started_at: Optional[datetime] = None
    auto_paused: bool = False
    hourly_rate_snapshot: Optional[int] = None


def format_duration(seconds: int) -> str:
    seconds = max(0, int(seconds or 0))
    return f"{seconds//3600:02}:{(seconds%3600)//60:02}:{seconds%60:02}"


def parse_datetime(value: str) -> datetime:
    return datetime.fromisoformat(value)


def to_qdatetime(value: datetime) -> QDateTime:
    return QDateTime(value.year, value.month, value.day, value.hour, value.minute, value.second)


def money(value) -> str:
    return "" if value is None else f"{int(value):,}"


class RecordEditDialog(QDialog):
    def __init__(self, db: DB, record_id: int, parent=None):
        super().__init__(parent)
        self.db = db
        self.record_id = record_id
        self.setWindowTitle("ویرایش رکورد")
        self.setModal(True)
        self.resize(620, 420)
        row = self.db.conn.execute("SELECT * FROM time_entries WHERE id=?", (record_id,)).fetchone()
        if not row:
            raise ValueError("رکورد پیدا نشد")
        self.row = row

        layout = QFormLayout(self)
        self.project = QComboBox()
        for p in self.db.conn.execute("SELECT * FROM projects ORDER BY name"):
            self.project.addItem(p["name"], dict(p))
            if p["id"] == row["project_id"]:
                self.project.setCurrentIndex(self.project.count() - 1)

        self.description = QTextEdit(row["task_description"] or "")
        self.start_time = QDateTimeEdit()
        self.start_time.setCalendarPopup(True)
        configure_persian_datetime_edit(self.start_time)
        self.start_time.setDateTime(to_qdatetime(parse_datetime(row["start_time"])))
        self.end_time = QDateTimeEdit()
        self.end_time.setCalendarPopup(True)
        configure_persian_datetime_edit(self.end_time)
        self.end_time.setDateTime(to_qdatetime(parse_datetime(row["end_time"])))
        self.rate = QLineEdit("" if row["hourly_rate_snapshot"] is None else str(row["hourly_rate_snapshot"]))
        self.amount = QLineEdit("" if row["amount"] is None else str(row["amount"]))
        self.duration_lbl = QLabel("")

        recalc = QPushButton("محاسبه مجدد")
        recalc.clicked.connect(self.recalculate)
        save = QPushButton("ذخیره")
        save.setObjectName("successButton")
        save.clicked.connect(self.save)
        cancel = QPushButton("لغو")
        cancel.clicked.connect(self.reject)
        buttons = QHBoxLayout()
        buttons.addWidget(recalc); buttons.addStretch(); buttons.addWidget(cancel); buttons.addWidget(save)

        layout.addRow("پروژه", self.project)
        layout.addRow("توضیحات", self.description)
        layout.addRow("شروع", self.start_time)
        layout.addRow("پایان", self.end_time)
        layout.addRow("نرخ ساعتی (ریال)", self.rate)
        layout.addRow("مبلغ (ریال)", self.amount)
        layout.addRow("مدت", self.duration_lbl)
        layout.addRow(buttons)
        layout.setFormAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignTop)
        layout.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        configure_rtl_widget(self)
        self.recalculate()

    def current_duration_seconds(self) -> int:
        start = self.start_time.dateTime().toPyDateTime()
        end = self.end_time.dateTime().toPyDateTime()
        return max(0, int((end - start).total_seconds()))

    def recalculate(self):
        dur = self.current_duration_seconds()
        self.duration_lbl.setText(format_duration(dur))
        rate_txt = self.rate.text().replace(",", "").strip()
        if rate_txt:
            self.amount.setText(str(int(dur / 3600 * int(rate_txt))))

    def save(self):
        project = self.project.currentData()
        if not project:
            QMessageBox.warning(self, "هشدار", "لطفاً یک پروژه انتخاب کنید.")
            return
        start = self.start_time.dateTime().toPyDateTime()
        end = self.end_time.dateTime().toPyDateTime()
        if end < start:
            QMessageBox.warning(self, "هشدار", "زمان پایان باید بعد از زمان شروع باشد.")
            return
        rate_txt = self.rate.text().replace(",", "").strip()
        amount_txt = self.amount.text().replace(",", "").strip()
        rate = int(rate_txt) if rate_txt else None
        amount = int(amount_txt) if amount_txt else None
        self.db.conn.execute(
            """
            UPDATE time_entries
            SET project_id=?, project_name_snapshot=?, task_description=?, start_time=?, end_time=?,
                duration_seconds=?, hourly_rate_snapshot=?, amount=?, updated_at=?
            WHERE id=?
            """,
            (
                project["id"], project["name"], self.description.toPlainText().strip(),
                start.isoformat(), end.isoformat(), self.current_duration_seconds(), rate, amount,
                datetime.now().isoformat(), self.record_id,
            ),
        )
        self.db.conn.commit()
        self.accept()


class SettingsDialog(QDialog):
    def __init__(self, db: DB, app_dir: Path, parent=None):
        super().__init__(parent)
        self.db = db
        self.app_dir = app_dir
        self.setWindowTitle("تنظیمات")
        self.resize(700, 500)
        self.setModal(True)
        tabs = QTabWidget()

        # General
        general = QWidget()
        gform = QFormLayout(general)
        self.hotkey = QLineEdit(self.db.get_setting("hotkey", "<shift>+q"))
        self.notify_enabled = QCheckBox()
        self.notify_enabled.setChecked(self.db.get_setting("notify_enabled", "1") == "1")
        self.notify_min = QSpinBox(); self.notify_min.setRange(5, 240); self.notify_min.setValue(int(self.db.get_setting("notify_minutes", "30")))
        self.idle_enabled = QCheckBox(); self.idle_enabled.setChecked(self.db.get_setting("idle_enabled", "1") == "1")
        self.idle_min = QSpinBox(); self.idle_min.setRange(1, 60); self.idle_min.setValue(int(self.db.get_setting("idle_minutes", "2")))
        self.default_rate = QLineEdit(self.db.get_setting("default_rate", ""))
        self.auto_start = QCheckBox()
        self.auto_start.setChecked(self.db.get_setting("auto_start", "0") == "1")
        self.show_on_startup = QCheckBox()
        self.show_on_startup.setChecked(self.db.get_setting("show_main_window_on_startup", "0") == "1")
        self.excel_dir = QLineEdit(self.db.get_setting("excel_dir", str(app_dir)))
        browse = QPushButton("انتخاب پوشه")
        browse.clicked.connect(self.pick_dir)
        h = QHBoxLayout(); h.addWidget(self.excel_dir); h.addWidget(browse)
        wrap = QWidget(); wrap.setLayout(h)
        gform.addRow("کلید میانبر سراسری", self.hotkey)
        gform.addRow("فعال‌سازی اعلان‌ها", self.notify_enabled)
        gform.addRow("یادآوری هر چند دقیقه", self.notify_min)
        gform.addRow("توقف خودکار هنگام بیکاری", self.idle_enabled)
        gform.addRow("دقیقه‌های بیکاری", self.idle_min)
        gform.addRow("نرخ ساعتی پیش‌فرض (ریال)", self.default_rate)
        gform.addRow("اجرای خودکار با ویندوز", self.auto_start)
        gform.addRow("نمایش پنجره اصلی هنگام اجرا", self.show_on_startup)
        gform.addRow("پوشه خروجی اکسل", wrap)

        # Projects
        projects = QWidget()
        pv = QVBoxLayout(projects)
        self.table = QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels(["شناسه", "نام", "نرخ (ریال)", "فعال", "گزارش"])
        self.table.setColumnHidden(0, True)
        pv.addWidget(self.table)
        hp = QHBoxLayout()
        add_btn, save_btn = QPushButton("افزودن"), QPushButton("ذخیره پروژه‌ها")
        hp.addWidget(add_btn); hp.addWidget(save_btn); hp.addStretch()
        pv.addLayout(hp)
        add_btn.clicked.connect(self.add_project_row)
        save_btn.clicked.connect(self.save_projects)

        tabs.addTab(general, "عمومی")
        tabs.addTab(projects, "پروژه‌ها")
        self.load_projects()

        btn = QPushButton("ذخیره تنظیمات")
        btn.setObjectName("primaryButton")
        btn.clicked.connect(self.save_settings)
        layout = QVBoxLayout(self); layout.addWidget(tabs); layout.addWidget(btn)
        gform.setFormAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignTop)
        gform.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        configure_rtl_widget(self)

    def pick_dir(self):
        path = QFileDialog.getExistingDirectory(self, "انتخاب پوشه", self.excel_dir.text())
        if path:
            self.excel_dir.setText(path)

    def load_projects(self):
        rows = self.db.conn.execute("SELECT * FROM projects ORDER BY id DESC").fetchall()
        self.table.setRowCount(len(rows))
        for i, r in enumerate(rows):
            self.table.setItem(i, 0, rtl_item(r["id"]))
            self.table.setItem(i, 1, rtl_item(r["name"]))
            self.table.setItem(i, 2, rtl_item("" if r["hourly_rate"] is None else str(r["hourly_rate"])))
            self.table.setCellWidget(i, 3, self.create_active_toggle(bool(r["is_active"])))
            report = QPushButton("مشاهده گزارش")
            report.clicked.connect(lambda _checked=False, project_id=r["id"]: self.open_project_report(project_id))
            self.table.setCellWidget(i, 4, report)

    def add_project_row(self):
        i = self.table.rowCount(); self.table.insertRow(i)
        self.table.setItem(i, 0, rtl_item(""))
        self.table.setItem(i, 1, rtl_item(""))
        self.table.setItem(i, 2, rtl_item(""))
        self.table.setCellWidget(i, 3, self.create_active_toggle(True))
        report = QPushButton("مشاهده گزارش")
        report.setEnabled(False)
        report.setToolTip("ابتدا پروژه را ذخیره کنید.")
        self.table.setCellWidget(i, 4, report)

    def save_projects(self):
        for r in range(self.table.rowCount()):
            pid = (self.table.item(r, 0).text() if self.table.item(r, 0) else "").strip()
            name = (self.table.item(r, 1).text() if self.table.item(r, 1) else "").strip()
            if not name:
                continue
            rate_txt = (self.table.item(r, 2).text() if self.table.item(r, 2) else "").replace(",", "").strip()
            rate = int(rate_txt) if rate_txt else None
            has_rate = 1 if rate is not None else 0
            active_widget = self.table.cellWidget(r, 3)
            active = 1 if isinstance(active_widget, QPushButton) and active_widget.isChecked() else 0
            now = datetime.now().isoformat()
            if pid:
                self.db.conn.execute("UPDATE projects SET name=?,hourly_rate=?,has_hourly_rate=?,is_active=?,updated_at=? WHERE id=?",
                                     (name, rate, has_rate, active, now, int(pid)))
            else:
                self.db.conn.execute("INSERT INTO projects(name,hourly_rate,has_hourly_rate,is_active,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                                     (name, rate, has_rate, active, now, now))
        self.db.conn.commit()
        self.load_projects()

    def create_active_toggle(self, checked: bool) -> QPushButton:
        button = QPushButton()
        button.setCheckable(True)
        button.setObjectName("activeToggle")
        button.setLayoutDirection(Qt.LayoutDirection.RightToLeft)

        def update_text(is_checked: bool):
            button.setText("روشن" if is_checked else "خاموش")

        button.toggled.connect(update_text)
        button.setChecked(checked)
        update_text(checked)
        return button

    def open_project_report(self, project_id: int):
        self.save_projects()
        parent = self.parent()
        if parent and hasattr(parent, "show_project_records"):
            parent.show_project_records(project_id)
            self.accept()

    def save_settings(self):
        self.save_projects()
        self.db.set_setting("hotkey", self.hotkey.text().strip() or "<shift>+q")
        self.db.set_setting("notify_enabled", int(self.notify_enabled.isChecked()))
        self.db.set_setting("notify_minutes", self.notify_min.value())
        self.db.set_setting("idle_enabled", int(self.idle_enabled.isChecked()))
        self.db.set_setting("idle_minutes", self.idle_min.value())
        self.db.set_setting("default_rate", self.default_rate.text().strip())
        self.db.set_setting("auto_start", int(self.auto_start.isChecked()))
        self.db.set_setting("show_main_window_on_startup", int(self.show_on_startup.isChecked()))
        self.db.set_setting("excel_dir", self.excel_dir.text().strip())

        try:
            MainWindow.sync_startup_shortcut(self.auto_start.isChecked())
        except Exception as e:
            logging.exception("auto-start sync failed")
            QMessageBox.warning(self, "اجرای خودکار", f"به‌روزرسانی اجرای خودکار ناموفق بود: {e}")
        self.accept()


class QuickStartPopup(QWidget):
    def __init__(self, main_window):
        super().__init__(main_window, Qt.WindowType.Popup | Qt.WindowType.FramelessWindowHint)
        self.main_window = main_window
        self.setObjectName("quickStartPopup")
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setWindowTitle("شروع سریع")
        self.setFixedWidth(340)

        root = QVBoxLayout(self)
        root.setContentsMargins(14, 12, 14, 14)
        root.setSpacing(10)

        header = QHBoxLayout()
        title = QLabel("شروع سریع")
        title.setObjectName("quickStartTitle")
        self.open_main_btn = QPushButton("↗")
        self.open_main_btn.setObjectName("quickOpenMainButton")
        self.open_main_btn.setToolTip("باز کردن پنجره اصلی")
        self.open_main_btn.setFixedSize(34, 30)
        self.open_main_btn.clicked.connect(self.open_main_window)
        header.addWidget(title)
        header.addStretch()
        header.addWidget(self.open_main_btn)

        self.project_combo = QComboBox()
        self.description = QTextEdit()
        self.description.setPlaceholderText("توضیحات فعالیت")
        self.description.setFixedHeight(74)
        self.status = QLabel("")
        self.status.setObjectName("quickStartStatus")
        self.start_btn = QPushButton("شروع")
        self.start_btn.setObjectName("primaryButton")
        self.start_btn.clicked.connect(self.start_timer)

        form = QFormLayout()
        form.setFormAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignTop)
        form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        form.addRow("پروژه", self.project_combo)
        form.addRow("توضیحات", self.description)

        root.addLayout(header)
        root.addLayout(form)
        root.addWidget(self.status)
        root.addWidget(self.start_btn)
        configure_rtl_widget(self)

    def load_projects(self):
        current = self.project_combo.currentData()
        self.project_combo.clear()
        selected = 0
        rows = self.main_window.db.conn.execute("SELECT * FROM projects WHERE is_active=1 ORDER BY name").fetchall()
        for row in rows:
            self.project_combo.addItem(row["name"], dict(row))
            if current == row["id"]:
                selected = self.project_combo.count() - 1
        if self.project_combo.count():
            self.project_combo.setCurrentIndex(selected)

        if self.main_window.timer_state:
            title = self.main_window.notification_activity_title()
            self.status.setText(f"تایمر فعال: {title}")
            self.start_btn.setEnabled(False)
            self.start_btn.setText("تایمر فعال است")
        elif self.project_combo.count() == 0:
            self.status.setText("ابتدا از تنظیمات یک پروژه فعال بسازید.")
            self.start_btn.setEnabled(False)
            self.start_btn.setText("شروع")
        else:
            self.status.setText("پروژه را انتخاب کنید، توضیحات را بنویسید و شروع را بزنید.")
            self.start_btn.setEnabled(True)
            self.start_btn.setText("شروع")

    def start_timer(self):
        data = self.project_combo.currentData()
        description = self.description.toPlainText().strip()
        if self.main_window.start_timer_from_quick_start(data, description):
            self.description.clear()
            self.close()

    def open_main_window(self):
        self.close()
        self.main_window.show_from_tray()

    def show_near_tray(self, tray: Optional[QSystemTrayIcon]):
        self.load_projects()
        self.adjustSize()
        tray_geometry = tray.geometry() if tray else None
        anchor = tray_geometry.center() if tray_geometry and tray_geometry.isValid() else QCursor.pos()
        screen = QApplication.screenAt(anchor) or QApplication.primaryScreen()
        available = screen.availableGeometry()
        x = anchor.x() - self.width() + 24
        if x < available.left():
            x = available.left() + 8
        if x + self.width() > available.right():
            x = available.right() - self.width() - 8
        if anchor.y() > available.center().y():
            y = anchor.y() - self.height() - 12
        else:
            y = anchor.y() + 12
        if y < available.top():
            y = available.top() + 8
        if y + self.height() > available.bottom():
            y = available.bottom() - self.height() - 8
        self.move(x, y)
        self.show()
        self.raise_()
        self.activateWindow()


class MainWindow(QMainWindow):
    @staticmethod
    def startup_shortcut_path() -> Path:
        startup_dir = Path(os.environ.get("APPDATA", str(Path.home()))) / "Microsoft" / "Windows" / "Start Menu" / "Programs" / "Startup"
        return startup_dir / f"{APP_NAME}.lnk"

    @staticmethod
    def current_app_path() -> Path:
        if getattr(sys, "frozen", False):
            return Path(sys.executable).resolve()
        return Path(__file__).resolve()

    @staticmethod
    def sync_startup_shortcut(enabled: bool):
        if os.name != "nt":
            return
        shortcut = MainWindow.startup_shortcut_path()
        shortcut.parent.mkdir(parents=True, exist_ok=True)

        if not enabled:
            if shortcut.exists():
                shortcut.unlink()
            return

        target = MainWindow.current_app_path()
        workdir = target.parent
        icon_location = str(target)
        escaped_shortcut = str(shortcut).replace("'", "''")
        escaped_target = str(target).replace("'", "''")
        escaped_workdir = str(workdir).replace("'", "''")
        escaped_icon = icon_location.replace("'", "''")
        ps_script = (
            f"$s=(New-Object -ComObject WScript.Shell).CreateShortcut('{escaped_shortcut}');"
            f"$s.TargetPath='{escaped_target}';"
            f"$s.WorkingDirectory='{escaped_workdir}';"
            f"$s.IconLocation='{escaped_icon}';"
            "$s.Save();"
        )
        subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps_script],
            check=True,
            capture_output=True,
            text=True,
        )
    def __init__(self, db, app_dir):
        super().__init__()
        self.db, self.app_dir = db, app_dir
        self.setWindowFlags(
            Qt.WindowType.Window
            | Qt.WindowType.Tool
            | Qt.WindowType.WindowMinMaxButtonsHint
            | Qt.WindowType.WindowCloseButtonHint
        )
        self.setWindowTitle("ردیاب زمان و درآمد")
        self.resize(980, 640)
        self.timer_state: Optional[TimerState] = None
        self.last_notification_at: Optional[datetime] = None
        self.hotkey_listener = None
        self.tray: Optional[QSystemTrayIcon] = None
        self.quick_start_popup: Optional[QuickStartPopup] = None
        self.is_quitting = False
        self.tray_menu = None
        self.open_act = None
        self.hide_act = None
        self.start_resume_act = None
        self.pause_act = None
        self.stop_act = None
        self.last_tick_at = datetime.now()

        self.init_ui()
        configure_rtl_widget(self)
        self.apply_dark()
        self.load_projects()
        self.restore_active_state()
        self.init_timers()
        self.init_tray()
        self.setup_hotkey()

    def init_ui(self):
        w = QWidget(); self.setCentralWidget(w)
        root = QVBoxLayout(w)
        root.setContentsMargins(22, 22, 22, 22)
        self.main_tabs = QTabWidget()
        root.addWidget(self.main_tabs)

        timer_tab = QWidget()
        v = QVBoxLayout(timer_tab)
        v.setSpacing(16)

        self.project_combo = QComboBox()
        self.desc = QTextEdit(); self.desc.setPlaceholderText("شرح فعالیت")
        self.timer_lbl = QLabel("00:00:00"); self.timer_lbl.setObjectName("timerLabel")
        self.amount_lbl = QLabel("0 ریال")
        self.amount_lbl.setObjectName("amountLabel")
        self.status_lbl = QLabel("وضعیت: آماده")
        self.status_lbl.setObjectName("statusLabel")
        self.banner = QLabel("")
        self.banner.setObjectName("bannerLabel")

        form = QFormLayout()
        form.setFormAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignTop)
        form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        form.setVerticalSpacing(12)
        form.addRow("پروژه", self.project_combo)
        form.addRow("توضیحات", self.desc)

        btns = QHBoxLayout()
        button_map = [
            ("شروع", self.start_timer, "primaryButton"),
            ("توقف موقت", self.pause_timer, ""),
            ("ادامه", self.resume_timer, ""),
            ("توقف و ذخیره", self.stop_save, "successButton"),
            ("لغو", self.cancel_timer, "dangerButton"),
            ("تنظیمات", self.open_settings, ""),
            ("خروجی اکسل", self.export_excel, ""),
        ]
        for t, fn, cls in button_map:
            b = QPushButton(t)
            if cls:
                b.setObjectName(cls)
            b.clicked.connect(fn)
            btns.addWidget(b)

        v.addLayout(form); v.addWidget(self.timer_lbl); v.addWidget(self.amount_lbl); v.addWidget(self.status_lbl); v.addWidget(self.banner); v.addLayout(btns)
        self.main_tabs.addTab(timer_tab, "تایمر")
        self.build_records_tab()

    def build_records_tab(self):
        records_tab = QWidget()
        layout = QVBoxLayout(records_tab)
        filters = QHBoxLayout()
        self.records_project = QComboBox(); self.records_project.addItem("همه پروژه‌ها", None)
        self.records_start = QDateEdit(); self.records_start.setCalendarPopup(True); configure_persian_date_edit(self.records_start)
        self.records_end = QDateEdit(); self.records_end.setCalendarPopup(True); configure_persian_date_edit(self.records_end)
        today = datetime.now().date()
        self.records_start.setDate(today.replace(day=1))
        self.records_end.setDate(today)
        refresh = QPushButton("به‌روزرسانی")
        refresh.clicked.connect(self.load_records)
        filters.addWidget(QLabel("پروژه")); filters.addWidget(self.records_project)
        filters.addWidget(QLabel("از")); filters.addWidget(self.records_start)
        filters.addWidget(QLabel("تا")); filters.addWidget(self.records_end)
        filters.addWidget(refresh); filters.addStretch()

        self.records_summary = QTableWidget(0, 4)
        self.records_summary.setHorizontalHeaderLabels(["پروژه", "تعداد رکورد", "جمع زمان", "جمع درآمد (ریال)"])
        self.records_summary.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)

        self.records_table = QTableWidget(0, 9)
        self.records_table.setHorizontalHeaderLabels(["شناسه", "تاریخ", "پروژه", "توضیحات", "شروع", "پایان", "مدت", "نرخ (ریال)", "مبلغ (ریال)"])
        self.records_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.records_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.records_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.records_table.doubleClicked.connect(self.edit_selected_record)

        actions = QHBoxLayout()
        edit = QPushButton("ویرایش انتخاب‌شده")
        delete = QPushButton("حذف انتخاب‌شده")
        delete.setObjectName("dangerButton")
        edit.clicked.connect(self.edit_selected_record)
        delete.clicked.connect(self.delete_selected_record)
        actions.addWidget(edit); actions.addWidget(delete); actions.addStretch()

        layout.addLayout(filters)
        layout.addWidget(QLabel("جمع پروژه‌ها"))
        layout.addWidget(self.records_summary)
        layout.addWidget(QLabel("رکوردها"))
        layout.addWidget(self.records_table)
        layout.addLayout(actions)
        self.main_tabs.addTab(records_tab, "رکوردها")

    def apply_dark(self):
        self.setStyleSheet("""
            QWidget { background:#0f172a; color:#e2e8f0; font-family:APP_FONT_STACK; font-size:14px; }
            QMainWindow { background:#0f172a; }
            QLineEdit, QTextEdit, QComboBox, QSpinBox, QDateEdit, QDateTimeEdit, QTableWidget, QTabWidget::pane {
                background:#1e293b; color:#f8fafc; border:1px solid #334155; border-radius:10px; padding:7px;
            }
            QComboBox::drop-down { border:none; }
            QPushButton {
                background:#334155; color:#f8fafc; border:1px solid #475569; border-radius:10px; padding:10px 14px; font-weight:600;
            }
            QPushButton:hover { background:#475569; }
            QPushButton#primaryButton { background:#2563eb; border:1px solid #3b82f6; }
            QPushButton#primaryButton:hover { background:#1d4ed8; }
            QPushButton#successButton { background:#059669; border:1px solid #10b981; }
            QPushButton#successButton:hover { background:#047857; }
            QPushButton#dangerButton { background:#dc2626; border:1px solid #ef4444; }
            QPushButton#dangerButton:hover { background:#b91c1c; }
            QPushButton#activeToggle { background:#475569; border:1px solid #64748b; padding:6px 10px; }
            QPushButton#activeToggle:checked { background:#059669; border:1px solid #10b981; }
            QLabel#timerLabel { font-size:52px; font-weight:800; color:#f8fafc; padding-top:8px; }
            QLabel#amountLabel { font-size:26px; font-weight:700; color:#93c5fd; }
            QLabel#statusLabel { color:#94a3b8; font-size:13px; }
            QLabel#bannerLabel { color:#facc15; font-weight:600; min-height:22px; }
            QWidget#quickStartPopup { border:1px solid #475569; border-radius:14px; }
            QLabel#quickStartTitle { font-size:16px; font-weight:800; color:#f8fafc; }
            QLabel#quickStartStatus { color:#cbd5e1; font-size:12px; }
            QPushButton#quickOpenMainButton { padding:2px; border-radius:8px; font-size:16px; font-weight:800; }
            QHeaderView::section { background:#1e293b; color:#cbd5e1; border:1px solid #334155; padding:6px; font-weight:700; }
            QTabBar::tab { background:#1e293b; color:#cbd5e1; padding:8px 14px; border-top-left-radius:8px; border-top-right-radius:8px; margin-right:4px; }
            QTabBar::tab:selected { background:#2563eb; color:#ffffff; }
        """.replace("APP_FONT_STACK", stylesheet_font_family()))

    def load_projects(self):
        self.project_combo.clear()
        rows = self.db.conn.execute("SELECT * FROM projects WHERE is_active=1 ORDER BY name").fetchall()
        for r in rows:
            self.project_combo.addItem(r["name"], dict(r))
        if hasattr(self, "records_project"):
            current = self.records_project.currentData()
            self.records_project.blockSignals(True)
            self.records_project.clear()
            self.records_project.addItem("همه پروژه‌ها", None)
            selected = 0
            for r in self.db.conn.execute("SELECT id,name FROM projects ORDER BY name"):
                self.records_project.addItem(r["name"], r["id"])
                if current == r["id"]:
                    selected = self.records_project.count() - 1
            self.records_project.setCurrentIndex(selected)
            self.records_project.blockSignals(False)
            self.load_records()

    def current_duration(self, at_time: Optional[datetime] = None):
        if not self.timer_state:
            return 0
        d = self.timer_state.accumulated_seconds
        if not self.timer_state.is_paused:
            d += int(((at_time or datetime.now()) - self.timer_state.start_time).total_seconds())
        return max(0, d)

    def begin_timer(self, data, description: str) -> bool:
        if self.timer_state:
            QMessageBox.information(self, "تایمر فعال", "یک تایمر از قبل فعال است. ابتدا آن را متوقف یا لغو کنید.")
            return False
        if not data:
            QMessageBox.warning(self, "هشدار", "ابتدا یک پروژه فعال بسازید.")
            return False
        self.timer_state = TimerState(
            project_id=data["id"], project_name=data["name"], task_description=description,
            start_time=datetime.now(), hourly_rate_snapshot=data["hourly_rate"] if data["has_hourly_rate"] else None
        )
        self.status_lbl.setText("وضعیت: در حال اجرا")
        self.save_active_state()
        self.update_tray_menu()
        return True

    def start_timer(self):
        return self.begin_timer(self.project_combo.currentData(), self.desc.toPlainText().strip())

    def start_timer_from_quick_start(self, data, description: str) -> bool:
        if not self.begin_timer(data, description):
            return False
        self.desc.setPlainText(description)
        if data:
            for index in range(self.project_combo.count()):
                item = self.project_combo.itemData(index)
                if item and item.get("id") == data["id"]:
                    self.project_combo.setCurrentIndex(index)
                    break
        return True

    def pause_timer(self, auto=False, effective_end: Optional[datetime] = None):
        if not self.timer_state or self.timer_state.is_paused:
            return
        self.timer_state.accumulated_seconds = self.current_duration(effective_end)
        self.timer_state.is_paused = True
        self.timer_state.pause_started_at = datetime.now()
        self.timer_state.auto_paused = auto
        self.status_lbl.setText("وضعیت: متوقف")
        self.save_active_state()

    def resume_timer(self):
        if not self.timer_state or not self.timer_state.is_paused:
            return
        self.timer_state.start_time = datetime.now()
        self.timer_state.is_paused = False
        self.timer_state.pause_started_at = None
        self.timer_state.auto_paused = False
        self.banner.setText("")
        self.status_lbl.setText("وضعیت: در حال اجرا")
        self.save_active_state()

    def stop_save(self):
        if not self.timer_state:
            return
        end = datetime.now()
        dur = self.current_duration()
        rate = self.timer_state.hourly_rate_snapshot
        amount = int(dur / 3600 * rate) if rate is not None else None
        now = datetime.now().isoformat()
        self.db.conn.execute("""INSERT INTO time_entries(project_id,project_name_snapshot,task_description,start_time,end_time,duration_seconds,hourly_rate_snapshot,amount,status,created_at,updated_at)
                           VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                           (self.timer_state.project_id, self.timer_state.project_name, self.timer_state.task_description,
                            self.timer_state.start_time.isoformat(), end.isoformat(), dur, rate, amount, "saved", now, now))
        self.db.conn.execute("DELETE FROM active_timer_state")
        self.db.conn.commit()
        self.timer_state = None
        self.status_lbl.setText("وضعیت: ذخیره شد")
        self.update_tray_menu()
        self.load_records()

    def cancel_timer(self):
        self.timer_state = None
        self.db.conn.execute("DELETE FROM active_timer_state")
        self.db.conn.commit()
        self.status_lbl.setText("وضعیت: لغو شد")
        self.timer_lbl.setText("00:00:00")
        self.amount_lbl.setText("0 ریال")
        self.update_tray_menu()

    def save_active_state(self, accumulated_override: Optional[int] = None, is_paused_override: Optional[bool] = None):
        if not self.timer_state:
            return
        accumulated = self.timer_state.accumulated_seconds if accumulated_override is None else accumulated_override
        is_paused = self.timer_state.is_paused if is_paused_override is None else is_paused_override
        self.db.conn.execute("DELETE FROM active_timer_state")
        self.db.conn.execute("""INSERT INTO active_timer_state(id,project_id,project_name_snapshot,task_description,start_time,accumulated_seconds,is_paused,pause_started_at,auto_paused,hourly_rate_snapshot,updated_at)
                           VALUES(1,?,?,?,?,?,?,?,?,?,?)""",
                           (self.timer_state.project_id, self.timer_state.project_name, self.timer_state.task_description,
                            self.timer_state.start_time.isoformat(), accumulated, int(is_paused),
                            self.timer_state.pause_started_at.isoformat() if self.timer_state.pause_started_at else None,
                            int(self.timer_state.auto_paused), self.timer_state.hourly_rate_snapshot, datetime.now().isoformat()))
        self.db.conn.commit()

    def restore_active_state(self):
        row = self.db.conn.execute("SELECT * FROM active_timer_state WHERE id=1").fetchone()
        if not row:
            return
        self.timer_state = TimerState(
            project_id=row["project_id"], project_name=row["project_name_snapshot"], task_description=row["task_description"] or "",
            start_time=datetime.fromisoformat(row["start_time"]), accumulated_seconds=row["accumulated_seconds"] or 0,
            is_paused=bool(row["is_paused"]), pause_started_at=datetime.fromisoformat(row["pause_started_at"]) if row["pause_started_at"] else None,
            auto_paused=bool(row["auto_paused"]), hourly_rate_snapshot=row["hourly_rate_snapshot"]
        )
        if not self.timer_state.is_paused:
            self.timer_state.is_paused = True
            self.timer_state.pause_started_at = datetime.now()
            self.timer_state.auto_paused = True
            self.banner.setText("تایمر بعد از اجرای دوباره برنامه متوقف نگه داشته شد تا زمان خاموشی/استندبای محاسبه نشود.")
            self.save_active_state()
        self.desc.setPlainText(self.timer_state.task_description)
        self.status_lbl.setText("وضعیت: متوقف" if self.timer_state.is_paused else "وضعیت: در حال اجرا")

    def init_timers(self):
        self.ui_timer = QTimer(self); self.ui_timer.setInterval(1000); self.ui_timer.timeout.connect(self.tick); self.ui_timer.start()
        self.idle_timer = QTimer(self); self.idle_timer.setInterval(5000); self.idle_timer.timeout.connect(self.check_idle); self.idle_timer.start()

    def tick(self):
        self.resume_after_activity_if_needed()
        self.pause_for_inactivity_if_needed()
        d = self.current_duration()
        self.timer_lbl.setText(f"{d//3600:02}:{(d%3600)//60:02}:{d%60:02}")
        if self.timer_state and self.timer_state.hourly_rate_snapshot:
            amount = int(d / 3600 * self.timer_state.hourly_rate_snapshot)
            self.amount_lbl.setText(f"{amount:,} ریال")
        else:
            self.amount_lbl.setText("0 ریال")
        self.notify_if_needed()
        if self.timer_state and not self.timer_state.is_paused:
            self.save_active_state(accumulated_override=d)
        self.last_tick_at = datetime.now()

    def notification_activity_title(self) -> str:
        if not self.timer_state:
            return ""
        description = self.timer_state.task_description.strip()
        if description:
            return f"{self.timer_state.project_name}، {description}"
        return self.timer_state.project_name

    def notify_if_needed(self):
        if not self.timer_state or self.timer_state.is_paused:
            return
        if self.db.get_setting("notify_enabled", "1") != "1":
            return
        mins = int(self.db.get_setting("notify_minutes", "30"))
        now = datetime.now()
        if self.last_notification_at and (now - self.last_notification_at).total_seconds() < mins * 60:
            return
        self.last_notification_at = now
        if self.tray:
            self.tray.showMessage("ردیاب زمان و درآمد", f"هنوز مشغول «{self.notification_activity_title()}» هستی؟", QSystemTrayIcon.MessageIcon.Information, 6000)

    def check_idle(self):
        self.resume_after_activity_if_needed()
        self.pause_for_inactivity_if_needed()

    def resume_after_activity_if_needed(self):
        if not self.timer_state or not self.timer_state.is_paused or not self.timer_state.auto_paused:
            return
        if self.db.get_setting("idle_enabled", "1") != "1":
            return
        if IdleMonitor.idle_seconds() <= AUTO_RESUME_ACTIVITY_SECONDS:
            self.resume_timer()
            self.banner.setText("فعالیت موس/کیبورد تشخیص داده شد؛ تایمر به صورت خودکار ادامه پیدا کرد.")

    def pause_for_inactivity_if_needed(self):
        if not self.timer_state or self.timer_state.is_paused:
            return
        if self.db.get_setting("idle_enabled", "1") != "1":
            return
        mins = int(self.db.get_setting("idle_minutes", "2"))
        now = datetime.now()
        gap_seconds = int((now - self.last_tick_at).total_seconds())
        if gap_seconds > max(30, mins * 60):
            self.pause_timer(auto=True, effective_end=self.last_tick_at)
            self.banner.setText("تایمر به دلیل وقفه طولانی سیستم (استندبای/خاموشی/قفل) متوقف شد و آن زمان محاسبه نشد.")
            return
        idle = IdleMonitor.idle_seconds()
        if idle >= mins * 60:
            last_input = now - timedelta(seconds=idle)
            self.pause_timer(auto=True, effective_end=max(self.timer_state.start_time, last_input))
            self.banner.setText("تایمر به دلیل عدم فعالیت موس/کیبورد متوقف شد و زمان بیکاری، استندبای یا خاموشی محاسبه نشد.")

    def build_app_icon(self) -> QIcon:
        icon = QIcon(self.db.get_setting("tray_icon_path", ""))
        if icon.isNull():
            icon = self.style().standardIcon(QStyle.StandardPixmap.SP_ComputerIcon)
        return icon

    def init_tray(self):
        self.tray = QSystemTrayIcon(self)
        app_icon = self.windowIcon()
        if app_icon.isNull():
            app_icon = self.build_app_icon()
            self.setWindowIcon(app_icon)
        self.tray.setIcon(app_icon)
        self.tray.setToolTip("ردیاب زمان و درآمد")
        self.tray_menu = QMenu()
        self.open_act = QAction("باز کردن / نمایش", self); self.open_act.triggered.connect(self.show_from_tray)
        self.hide_act = QAction("مخفی کردن", self); self.hide_act.triggered.connect(self.hide)
        self.start_resume_act = QAction("شروع تایمر", self); self.start_resume_act.triggered.connect(self.start_or_resume_timer)
        self.pause_act = QAction("توقف موقت تایمر", self); self.pause_act.triggered.connect(self.pause_timer)
        self.stop_act = QAction("توقف و ذخیره", self); self.stop_act.triggered.connect(self.stop_save)
        settings_act = QAction("تنظیمات", self); settings_act.triggered.connect(self.open_settings)
        exit_act = QAction("خروج", self); exit_act.triggered.connect(self.quit_app)
        for a in [self.open_act, self.hide_act, self.start_resume_act, self.pause_act, self.stop_act, settings_act, exit_act]:
            self.tray_menu.addAction(a)
        self.tray_menu.aboutToShow.connect(self.update_tray_menu)
        self.tray.setContextMenu(self.tray_menu)
        self.tray.activated.connect(self.on_tray_activated)
        self.tray.show()
        self.update_tray_menu()

    def setup_hotkey(self):
        try:
            from pynput import keyboard
            hotkey = self.db.get_setting("hotkey", "<shift>+q")
            self.hotkey_listener = keyboard.GlobalHotKeys({hotkey: self.toggle_window})
            self.hotkey_listener.start()
        except Exception as e:
            logging.exception("hotkey failed")
            self.status_lbl.setText(f"وضعیت: خطای کلید میانبر ({e})")

    def toggle_window(self):
        if self.isHidden():
            self.show_from_tray()
        else:
            self.hide()

    def show_quick_start_popup(self):
        if self.quick_start_popup is None:
            self.quick_start_popup = QuickStartPopup(self)
        self.quick_start_popup.show_near_tray(self.tray)

    def show_from_tray(self):
        if self.quick_start_popup:
            self.quick_start_popup.close()
        self.showMaximized()
        self.raise_()
        self.activateWindow()

    def on_tray_activated(self, reason):
        if reason == QSystemTrayIcon.ActivationReason.Trigger:
            self.show_quick_start_popup()
        elif reason == QSystemTrayIcon.ActivationReason.DoubleClick:
            self.show_from_tray()

    def start_or_resume_timer(self):
        if self.timer_state and self.timer_state.is_paused:
            self.resume_timer()
        elif not self.timer_state:
            self.start_timer()

    def update_tray_menu(self):
        if self.timer_state:
            if self.timer_state.is_paused:
                self.start_resume_act.setText("ادامه تایمر")
                self.start_resume_act.setEnabled(True)
                self.pause_act.setEnabled(False)
                self.stop_act.setEnabled(True)
            else:
                self.start_resume_act.setText("تایمر در حال اجراست")
                self.start_resume_act.setEnabled(False)
                self.pause_act.setEnabled(True)
                self.stop_act.setEnabled(True)
        else:
            self.start_resume_act.setText("شروع تایمر")
            self.start_resume_act.setEnabled(True)
            self.pause_act.setEnabled(False)
            self.stop_act.setEnabled(False)

    def records_filter(self):
        start = self.records_start.date().toPyDate()
        end = self.records_end.date().toPyDate()
        args = [datetime.combine(start, datetime.min.time()).isoformat(), (datetime.combine(end, datetime.min.time()) + timedelta(days=1)).isoformat()]
        where = "WHERE status='saved' AND start_time>=? AND start_time<?"
        pid = self.records_project.currentData()
        if pid:
            where += " AND project_id=?"
            args.append(pid)
        return where, args

    def load_records(self):
        if not hasattr(self, "records_table"):
            return
        where, args = self.records_filter()
        rows = self.db.conn.execute(f"SELECT * FROM time_entries {where} ORDER BY start_time DESC", args).fetchall()
        self.records_table.setRowCount(len(rows))
        for i, r in enumerate(rows):
            st = parse_datetime(r["start_time"])
            en = parse_datetime(r["end_time"])
            values = [
                r["id"], format_jalali_date(st), r["project_name_snapshot"], r["task_description"] or "",
                st.strftime("%H:%M:%S"), en.strftime("%H:%M:%S"), format_duration(r["duration_seconds"]),
                money(r["hourly_rate_snapshot"]), money(r["amount"]),
            ]
            for col, value in enumerate(values):
                item = rtl_item(value)
                if col == 0:
                    item.setData(Qt.ItemDataRole.UserRole, r["id"])
                self.records_table.setItem(i, col, item)

        summary = self.db.conn.execute(
            f"""
            SELECT project_name_snapshot, COUNT(*) AS count_records,
                   COALESCE(SUM(duration_seconds),0) AS total_seconds,
                   COALESCE(SUM(amount),0) AS total_amount
            FROM time_entries {where}
            GROUP BY project_name_snapshot
            ORDER BY project_name_snapshot
            """,
            args,
        ).fetchall()
        self.records_summary.setRowCount(len(summary) + 1)
        total_count = total_seconds = total_amount = 0
        for i, r in enumerate(summary):
            total_count += r["count_records"] or 0
            total_seconds += r["total_seconds"] or 0
            total_amount += r["total_amount"] or 0
            for col, value in enumerate([r["project_name_snapshot"], r["count_records"], format_duration(r["total_seconds"]), money(r["total_amount"])]):
                self.records_summary.setItem(i, col, rtl_item(value))
        last = len(summary)
        for col, value in enumerate(["جمع کل", total_count, format_duration(total_seconds), money(total_amount)]):
            self.records_summary.setItem(last, col, rtl_item(value))

    def show_project_records(self, project_id: int):
        if not hasattr(self, "records_project"):
            return
        for index in range(self.records_project.count()):
            if self.records_project.itemData(index) == project_id:
                self.records_project.setCurrentIndex(index)
                break
        records_index = self.main_tabs.indexOf(self.records_table.parentWidget())
        if records_index >= 0:
            self.main_tabs.setCurrentIndex(records_index)
        self.load_records()
        self.show_from_tray()

    def selected_record_id(self):
        row = self.records_table.currentRow()
        if row < 0:
            return None
        item = self.records_table.item(row, 0)
        return int(item.data(Qt.ItemDataRole.UserRole) or item.text()) if item else None

    def edit_selected_record(self):
        record_id = self.selected_record_id()
        if not record_id:
            QMessageBox.information(self, "رکوردها", "لطفاً ابتدا یک رکورد انتخاب کنید.")
            return
        dlg = RecordEditDialog(self.db, record_id, self)
        if dlg.exec():
            self.load_records()

    def delete_selected_record(self):
        record_id = self.selected_record_id()
        if not record_id:
            QMessageBox.information(self, "رکوردها", "لطفاً ابتدا یک رکورد انتخاب کنید.")
            return
        r = QMessageBox.question(
            self, "حذف رکورد", "آیا از حذف این رکورد مطمئن هستید؟",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if r != QMessageBox.StandardButton.Yes:
            return
        self.db.conn.execute("DELETE FROM time_entries WHERE id=?", (record_id,))
        self.db.conn.commit()
        self.load_records()

    def open_settings(self):
        dlg = SettingsDialog(self.db, self.app_dir, self)
        if dlg.exec():
            self.load_projects()

    def export_excel(self):
        dlg = QDialog(self); dlg.setWindowTitle("گزارش اکسل")
        layout = QFormLayout(dlg)
        period = QComboBox()
        for text, key in [("امروز", "today"), ("این هفته", "this_week"), ("این ماه", "this_month"), ("ماه قبل", "last_month"), ("دلخواه", "custom")]:
            period.addItem(text, key)
        project = QComboBox(); project.addItem("همه", None)
        for r in self.db.conn.execute("SELECT id,name FROM projects ORDER BY name"): project.addItem(r["name"], r["id"])
        start = QDateEdit(); start.setCalendarPopup(True); configure_persian_date_edit(start); start.setDate(datetime.now().date())
        end = QDateEdit(); end.setCalendarPopup(True); configure_persian_date_edit(end); end.setDate(datetime.now().date())
        go = QPushButton("خروجی گرفتن")
        layout.addRow("بازه", period); layout.addRow("پروژه", project); layout.addRow("شروع", start); layout.addRow("پایان", end); layout.addRow(go)
        layout.setFormAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignTop)
        layout.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        configure_rtl_widget(dlg)

        def do_export():
            sdt, edt = self.resolve_period(period.currentData(), start.date().toPyDate(), end.date().toPyDate())
            pid = project.currentData()
            q = "SELECT * FROM time_entries WHERE status='saved' AND start_time>=? AND start_time<?"
            args = [sdt.isoformat(), (edt + timedelta(days=1)).isoformat()]
            if pid:
                q += " AND project_id=?"; args.append(pid)
            rows = self.db.conn.execute(q, args).fetchall()
            wb = Workbook(); ws = wb.active; ws.title = "گزارش"
            headers = ["ردیف", "تاریخ", "نام پروژه", "شرح فعالیت", "ساعت شروع", "ساعت پایان", "مدت", "نرخ ساعتی (ریال)", "مبلغ (ریال)"]
            ws.append(headers)
            for c in ws[1]:
                c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="333333")
            total_sec = total_amount = 0
            for idx, r in enumerate(rows, 1):
                st = datetime.fromisoformat(r["start_time"]); en = datetime.fromisoformat(r["end_time"])
                dur = r["duration_seconds"] or 0
                total_sec += dur; total_amount += r["amount"] or 0
                ws.append([idx, format_jalali_date(st), r["project_name_snapshot"], r["task_description"], st.strftime("%H:%M"), en.strftime("%H:%M"), f"{dur//3600:02}:{(dur%3600)//60:02}:{dur%60:02}", r["hourly_rate_snapshot"], r["amount"]])
            ws.append([]); ws.append(["", "", "", "", "", "جمع کل زمان", f"{total_sec//3600:02}:{(total_sec%3600)//60:02}:{total_sec%60:02}", "جمع کل مبلغ (ریال)", total_amount])
            for i in range(1, 10): ws.column_dimensions[get_column_letter(i)].width = 18
            out_dir = Path(self.db.get_setting("excel_dir", str(self.app_dir))); out_dir.mkdir(parents=True, exist_ok=True)
            file = out_dir / f"گزارش-زمان-درآمد-{format_jalali_date(datetime.now())}.xlsx"
            wb.save(file)
            QMessageBox.information(self, "انجام شد", f"ذخیره شد: {file}")
            dlg.accept()

        go.clicked.connect(do_export)
        dlg.exec()

    def resolve_period(self, p, custom_s, custom_e):
        now = datetime.now().date()
        if p == "today":
            return now, now
        if p == "this_week":
            days_since_saturday = (now.weekday() + 2) % 7
            return now - timedelta(days=days_since_saturday), now
        if p == "this_month":
            start, _end = jalali_month_range(now)
            return start, now
        if p == "last_month":
            this_start, _this_end = jalali_month_range(now)
            previous_day = this_start - timedelta(days=1)
            return jalali_month_range(previous_day)
        return custom_s, custom_e

    def quit_app(self):
        if self.timer_state and not self.timer_state.is_paused:
            r = QMessageBox.question(
                self,
                "خروج",
                "یک تایمر فعال وجود دارد. آیا از خروج مطمئن هستید؟",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            )
            if r != QMessageBox.StandardButton.Yes:
                return
        self.is_quitting = True
        if self.tray:
            self.tray.hide()
        QApplication.quit()

    def closeEvent(self, event: QCloseEvent):
        if self.is_quitting:
            event.accept()
            return
        event.ignore()
        self.hide()


def main():
    app = QApplication(sys.argv)
    apply_app_font(app)
    if SingleInstanceServer.signal_existing():
        QMessageBox.information(None, "ردیاب زمان و درآمد", "برنامه از قبل باز است و همان پنجره فعال شد.")
        return

    d = app_data_dir()
    setup_logging(d)
    db = DB(d / "time_income.db")
    win = MainWindow(db, d)
    single_instance = SingleInstanceServer(win.show_from_tray)
    if not single_instance.listen():
        QMessageBox.warning(None, "ردیاب زمان و درآمد", "برنامه نتوانست قفل اجرای تکی را بسازد.")
        return
    app.setWindowIcon(win.windowIcon())
    if db.get_setting("show_main_window_on_startup", "0") == "1":
        win.show_from_tray()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
